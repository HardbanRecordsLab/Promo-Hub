"""Builds database/global-database.xlsx from database/global-database.csv.

Run after `npm run export:db`:  python3 scripts/export-database-xlsx.py
Sheet "Database": every entity with readable labels, filters and clickable links.
Sheet "Summary": live COUNTIF formulas by region, type, creator type and country.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(__file__), '..')
SRC = os.path.join(ROOT, 'database', 'global-database.csv')
OUT = os.path.join(ROOT, 'database', 'global-database.xlsx')

COUNTRIES = {
    'AE': 'United Arab Emirates', 'AR': 'Argentina', 'AT': 'Austria', 'AU': 'Australia', 'BE': 'Belgium',
    'BR': 'Brazil', 'CA': 'Canada', 'CH': 'Switzerland', 'CO': 'Colombia', 'CZ': 'Czechia', 'DE': 'Germany',
    'DK': 'Denmark', 'EE': 'Estonia', 'ES': 'Spain', 'FI': 'Finland', 'FR': 'France', 'GB': 'United Kingdom',
    'HK': 'Hong Kong', 'HU': 'Hungary', 'IE': 'Ireland', 'IL': 'Israel', 'IN': 'India', 'IS': 'Iceland',
    'IT': 'Italy', 'JP': 'Japan', 'KR': 'South Korea', 'LT': 'Lithuania', 'LV': 'Latvia', 'MX': 'Mexico',
    'NG': 'Nigeria', 'NL': 'Netherlands', 'NO': 'Norway', 'NZ': 'New Zealand', 'PL': 'Poland', 'PS': 'Palestine',
    'PT': 'Portugal', 'RS': 'Serbia', 'SA': 'Saudi Arabia', 'SE': 'Sweden', 'SG': 'Singapore', 'SI': 'Slovenia',
    'SK': 'Slovakia', 'UG': 'Uganda', 'US': 'United States', 'WW': 'Worldwide', 'ZA': 'South Africa',
}
REGIONS = {'EU': 'Europe', 'NA': 'North America', 'LA': 'Latin America', 'AF': 'Africa', 'AS': 'Asia',
           'ME': 'Middle East', 'OC': 'Oceania', 'WW': 'Worldwide'}
KINDS = {
    'playlist': 'Playlist', 'radio': 'FM / DAB radio', 'webradio': 'Online radio', 'college': 'College radio',
    'podcast': 'Podcast', 'youtube': 'YouTube / sessions', 'tiktok': 'TikTok creator', 'instagram': 'Instagram',
    'blog': 'Blog', 'newsletter': 'Newsletter', 'press': 'Press / media', 'magazine': 'Magazine',
    'community': 'Community', 'brand': 'Brand', 'dj': 'Record pool', 'club': 'Live / venue series',
    'twitch': 'Twitch', 'tv': 'TV', 'sync': 'Sync / licensing', 'mixseries': 'Mix series', 'promoter': 'Promoter',
    'publisher': 'Publisher', 'audiobook': 'E-book / audio / self-pub', 'bookstore': 'Bookshop',
    'festival': 'Festival / fair / showcase', 'label': 'Record label', 'platform': 'Platform / tool',
    'distributor': 'Distributor', 'rights': 'Rights organisation', 'funding': 'Funding / export office',
    'other': 'Other (prizes, contests)',
}
SUBMISSION = {'open': 'Open submissions', 'contact': 'Editorial / A&R contact', 'platform': 'Via a platform',
              'paid': 'Paid service', 'invite': 'Curated / invitation'}
MODELS = {'free': 'Free', 'commission': 'Commission', 'revshare': 'Revenue share', 'fee': 'Review fee', 'barter': 'Barter'}
CREATORS = {'music': 'Musician', 'dj': 'DJ / producer', 'writer': 'Writer'}

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', start_color='1F2430')
HEAD_FONT = Font(name=FONT, bold=True, color='FFFFFF')
BODY_FONT = Font(name=FONT, size=10)
LINK_FONT = Font(name=FONT, size=10, color='0563C1', underline='single')

with open(SRC, encoding='utf-8-sig', newline='') as f:
    rows = list(csv.DictReader(f))

wb = Workbook()
ws = wb.active
ws.title = 'Database'
headers = ['Name', 'Type', 'Creators', 'Country', 'Country code', 'Region', 'Website', 'How to reach',
           'Proposed model', 'Recommended level (1–5)', 'Note (EN)', 'Note (PL)', 'Stage', 'Verified']
ws.append(headers)
for r in rows:
    ws.append([
        r['name'], KINDS.get(r['kind'], r['kind']),
        ', '.join(CREATORS.get(c, c) for c in r['creatorTypes'].split('|') if c),
        COUNTRIES.get(r['country'], r['country']), r['country'], REGIONS.get(r['region'], r['region']),
        r['url'], SUBMISSION.get(r['submission'], r['submission']), MODELS.get(r['model'], r['model']),
        int(r['tier'] or 0), r['notes'], r['notesPl'], r['stage'], 'yes' if r['verified'] == 'true' else 'no',
    ])

for cell in ws[1]:
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 30
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical='top', wrap_text=cell.column in (11, 12))
    link = row[6]
    link.hyperlink = link.value
    link.font = LINK_FONT
widths = [34, 24, 26, 18, 9, 16, 38, 24, 16, 12, 48, 48, 11, 9]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'B2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'

# ---------- Summary (live formulas over the Database sheet) ----------
last = ws.max_row
sm = wb.create_sheet('Summary')
sm['A1'] = 'Promo Hub – Global Entity Database'
sm['A1'].font = Font(name=FONT, bold=True, size=14)
sm['A2'] = ('All entries are PROSPECTS (outreach targets), not partners. Submission terms are based on public '
            'information and may have changed – verify on the official website before contacting or publishing.')
sm['A2'].font = Font(name=FONT, italic=True, size=10, color='7F7F7F')
sm['A4'] = 'Total entities'
sm['B4'] = f'=COUNTA(Database!A2:A{last})'
sm['A5'] = 'Countries (incl. Worldwide)'
sm['B5'] = f'=SUMPRODUCT(1/COUNTIF(Database!E2:E{last},Database!E2:E{last}))'
for c in ('A4', 'A5'):
    sm[c].font = Font(name=FONT, bold=True)
for c in ('B4', 'B5'):
    sm[c].font = Font(name=FONT, bold=True)


def block(start_row, title, labels, column, wildcard=False):
    sm.cell(row=start_row, column=1, value=title).font = Font(name=FONT, bold=True, color='FFFFFF')
    sm.cell(row=start_row, column=1).fill = HEAD_FILL
    sm.cell(row=start_row, column=2, value='Count').font = Font(name=FONT, bold=True, color='FFFFFF')
    sm.cell(row=start_row, column=2).fill = HEAD_FILL
    r = start_row + 1
    for label in labels:
        sm.cell(row=r, column=1, value=label).font = BODY_FONT
        crit = f'"*{label}*"' if wildcard else f'A{r}'
        sm.cell(row=r, column=2, value=f'=COUNTIF(Database!{column}2:{column}{last},{crit})').font = BODY_FONT
        r += 1
    return r + 1


row = block(7, 'By creator type', list(CREATORS.values()), 'C', wildcard=True)
row = block(row, 'By region', list(REGIONS.values()), 'F')
used_kinds = sorted({KINDS.get(r['kind'], r['kind']) for r in rows})
row = block(row, 'By type', used_kinds, 'B')
row = block(row, 'By how to reach', list(SUBMISSION.values()), 'H')
row = block(row, 'By proposed model', list(MODELS.values()), 'I')
used_countries = sorted({COUNTRIES.get(r['country'], r['country']) for r in rows})
block(row, 'By country', used_countries, 'D')
sm.column_dimensions['A'].width = 34
sm.column_dimensions['B'].width = 12
sm.sheet_view.showGridLines = False
wb.move_sheet('Summary', offset=-1)
wb.active = 0

wb.save(OUT)
print(f'Saved {len(rows)} rows -> {os.path.relpath(OUT)}')
